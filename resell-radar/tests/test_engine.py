import tempfile
import unittest
from dataclasses import replace

from resell_radar import presets
from resell_radar.config import DEFAULT_REPAIR_COSTS, Settings
from resell_radar.engine import Engine
from tests.helpers import FakeBrowser, card, item_page, search_page

CLEAN = "Состояние отличное, АКБ 93%, полный комплект, коробка и чек в наличии, без дефектов"
REBOOT = "Телефон периодически перезагружается сам, цена такая из-за этого, остальное работает"
SCREEN = "Разбит экран, остальное работает отлично, АКБ 90%, комплект полный, торг уместен"

BASE_PRICES = [20000, 20500, 19800, 21000, 20200, 19500, 20800, 20100]


def make_search(overrides=None):
    prices = dict(overrides or {})
    cards = [card(100 + i, "Samsung Galaxy S22, 8/128 ГБ", p) for i, p in enumerate(BASE_PRICES)]
    cards += [
        card(901, "Samsung Galaxy S22, 8/128 ГБ", 14000),   # выгодный и целый
        card(902, "Samsung Galaxy S22, 8/128 ГБ", 14500),   # перезагружается
        card(903, "Samsung Galaxy S22, 8/128 ГБ", 9000),    # разбит экран
        card(904, "Samsung Galaxy S22, 8/128 ГБ", 25000),   # дороже лимита
        card(905, "Samsung Galaxy S22, 8/128 ГБ", prices.get(905, 19000)),  # маржа мала
        card(906, "Чехол Samsung Galaxy S22", 600),         # аксессуар
        card(907, "Samsung Galaxy S22 Ultra, 12/256 ГБ", 15000),  # другая модель
    ]
    return search_page(cards)


ITEMS = {901: item_page(CLEAN), 902: item_page(REBOOT), 903: item_page(SCREEN),
         905: item_page(CLEAN)}


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.item = replace(presets.all_presets()[0])   # Samsung Galaxy S22
        self.settings = Settings(data_dir=self.tmp, max_buy=18000, min_margin=1500,
                                 items=[self.item], min_group_size=3,
                                 sell_basis="median", sell_discount_pct=5.0,
                                 repair_costs={**DEFAULT_REPAIR_COSTS, "screen": 5000})

    def engine(self, browser, settings=None):
        return Engine(settings or self.settings, browser, sleep=lambda s: None)

    @staticmethod
    def ids(eng):
        return [r["id"] for r in eng.results.rows()]

    def browser(self, search=None, **kw):
        return FakeBrowser(search={"q=Samsung+Galaxy+S22": search or make_search()},
                           items={k: v for k, v in ITEMS.items()}, **kw)

    def test_only_working_sends_only_clean_phone(self):
        b = self.browser()
        eng = self.engine(b)
        self.assertFalse(eng.run_cycle())
        self.assertEqual(self.ids(eng), ["901"])
        row = eng.results.rows()[0]
        self.assertTrue(row["url"].endswith("phone_901"))
        self.assertTrue(int(row["margin"]) >= 1500)
        self.assertIn("risk", row)
        # лишние объявления даже не открывались
        for skipped in ("904", "905", "906", "907"):
            self.assertFalse(any(u.endswith(f"phone_{skipped}") for u in b.opened), skipped)
        self.assertEqual(len(eng.results.rows()), 1)

    def test_repairer_gets_screen_phone_with_repair_cost(self):
        s = replace(self.settings, repairable=["screen"])
        eng = self.engine(self.browser(), s)
        eng.run_cycle()
        self.assertEqual(sorted(self.ids(eng)), ["901", "903"])   # перезагрузки (902) он чинить не умеет
        rows = {r["id"]: r for r in eng.results.rows()}
        self.assertEqual(rows["903"]["repair_cost"], "5000")
        sell = int(rows["903"]["sell_price"])
        self.assertTrue(18500 <= sell <= 19500)          # ~95% от медианы около 20 000
        self.assertEqual(rows["903"]["margin"], str(round(sell * 0.89 - 9000 - 5000)))

    def test_no_duplicates_on_second_cycle(self):
        eng = self.engine(self.browser())
        eng.run_cycle()
        first = len(eng.results.rows())
        eng.run_cycle()
        self.assertEqual(len(eng.results.rows()), first)

    def test_price_drop_triggers_new_review(self):
        eng = self.engine(self.browser())
        eng.run_cycle()
        self.assertEqual(self.ids(eng), ["901"])
        eng.browser = self.browser(search=make_search({905: 14000}))
        eng.run_cycle()
        self.assertEqual(self.ids(eng), ["901", "905"])

    def test_blocked_on_search(self):
        eng = self.engine(self.browser(blocked=["q=Samsung"]))
        self.assertTrue(eng.run_cycle())
        self.assertEqual(eng.results.rows(), [])

    def test_blocked_on_item_page_is_retried_later(self):
        eng = self.engine(self.browser(blocked=["phone_901"]))
        self.assertTrue(eng.run_cycle())
        self.assertTrue(eng.seen.needs_review("901", 14000))   # не помечено как разобранное

    def test_unknown_description_is_flagged(self):
        b = FakeBrowser(search={"q=Samsung+Galaxy+S22": make_search()}, items={})  # страницы не открываются
        eng = self.engine(b)
        eng.run_cycle()
        rows = eng.results.rows()
        self.assertTrue(rows)
        self.assertTrue(all("описание не удалось проверить" in r["risk_reasons"] for r in rows))

    def test_state_survives_restart(self):
        self.engine(self.browser()).run_cycle()
        eng2 = self.engine(self.browser())
        eng2.run_cycle()
        self.assertEqual(self.ids(eng2), ["901"])   # новый объект, те же файлы: повторов нет

    def test_search_url_has_private_sellers_filter_by_default(self):
        eng = self.engine(self.browser())
        self.assertIn("user=1", eng.search_url(self.item))
        any_sellers = self.engine(self.browser(), replace(self.settings, seller="any"))
        self.assertNotIn("user=", any_sellers.search_url(self.item))

    def test_user_url_is_used_as_is(self):
        url = "https://www.avito.ru/ekaterinburg/telefony?q=pixel&user=2&pmin=5000&pmax=9000"
        item = replace(self.item, url=url)
        got = self.engine(self.browser()).search_url(item)
        self.assertIn("user=2", got)             # выбор пользователя не переопределён
        self.assertNotIn("user=1", got)
        self.assertNotIn("pmax", got)             # иначе рынок считался бы только по дешёвым

    def test_whole_model_variant_for_non_phones(self):
        item = replace(self.item, variant="none", name="AirPods Pro")
        eng = self.engine(self.browser(), replace(self.settings, items=[item]))
        eng.run_cycle()
        self.assertTrue(any(g.endswith("| все") for g in eng.market.data))

    def test_favorites_added_when_enabled(self):
        b = self.browser()
        eng = self.engine(b, replace(self.settings, favorite=True))
        eng.run_cycle()
        self.assertEqual(len(b.favorites), 1)
        self.assertTrue(b.favorites[0].endswith("phone_901"))
        self.assertEqual(self.ids(eng), ["901"])

    def test_favorites_not_used_when_disabled(self):
        b = self.browser()
        self.engine(b).run_cycle()
        self.assertEqual(b.favorites, [])

    def test_favorite_failure_still_saves_result(self):
        b = self.browser()
        b.favorite_status = "unconfirmed"
        eng = self.engine(b, replace(self.settings, favorite=True, repairable=["screen"]))
        eng.run_cycle()
        self.assertEqual(sorted(self.ids(eng)), ["901", "903"])

    def test_excel_table_is_written(self):
        from openpyxl import load_workbook
        eng = self.engine(self.browser(), replace(self.settings, repairable=["screen"]))
        eng.run_cycle()
        ws = load_workbook(eng.results.xlsx_path).active
        self.assertEqual(ws.cell(1, 1).value, "Дата")
        self.assertEqual(ws.max_row, 3)                               # заголовок + 2 находки
        margins = [ws.cell(r, 9).value for r in (2, 3)]
        self.assertEqual(margins, sorted(margins, reverse=True))      # лучшие сверху
        self.assertTrue(ws.cell(2, 15).hyperlink.target.endswith(("phone_901", "phone_903")))

    def test_blocked_on_favorite_is_retried(self):
        class Flaky(FakeBrowser):
            def add_favorite(self, url):
                from resell_radar.browser import Blocked
                raise Blocked("капча")
        b = Flaky(search={"q=Samsung+Galaxy+S22": make_search()}, items=dict(ITEMS))
        eng = self.engine(b, replace(self.settings, favorite=True))
        self.assertTrue(eng.run_cycle())
        self.assertTrue(eng.seen.needs_review("901", 14000))

    def test_example_listing_drives_the_search(self):
        example = "https://www.avito.ru/ekaterinburg/telefony/phone_4012345678"
        item = replace(self.item, name="Мой S22", query="", match="", exclude="", example=example)
        pages = {"4012345678": item_page("описание", title="Samsung Galaxy S22, 8/128 ГБ, чёрный")}
        pages.update(ITEMS)
        b = FakeBrowser(search={"q=samsung+galaxy+s22": make_search()}, items=pages)
        eng = self.engine(b, replace(self.settings, items=[item]))
        eng.run_cycle()
        self.assertEqual(self.ids(eng), ["901"])                       # S22 Ultra (907) отсеян по образцу
        self.assertTrue(any("q=samsung+galaxy+s22" in u for u in b.opened))
        # заголовок со страницы запомнен, второй раз образец не открывается
        b2 = FakeBrowser(search={"q=samsung+galaxy+s22": make_search()}, items=pages)
        self.engine(b2, replace(self.settings, items=[item])).run_cycle()
        self.assertFalse(any(u == example for u in b2.opened))

    def test_example_falls_back_to_url_slug(self):
        example = "https://www.avito.ru/ekaterinburg/telefony/samsung_galaxy_s22_8_128_gb_4012345678"
        item = replace(self.item, name="Мой S22", query="", match="", exclude="", example=example)
        b = FakeBrowser(search={"q=samsung+galaxy+s22": make_search()}, items=dict(ITEMS))   # страница образца не открылась
        eng = self.engine(b, replace(self.settings, items=[item]))
        eng.run_cycle()
        self.assertTrue(any("q=samsung+galaxy+s22" in u for u in b.opened))
        self.assertFalse(eng._examples)                                # по слагу не кэшируем

    def test_search_url(self):
        eng = self.engine(self.browser(), replace(self.settings, region="ekaterinburg"))
        url = eng.search_url(self.item)
        self.assertTrue(url.startswith("https://www.avito.ru/ekaterinburg/telefony?q=Samsung+Galaxy+S22"))
        self.assertNotIn("pmax", url)


if __name__ == "__main__":
    unittest.main()
