"""Что уже видели и что нашли: JSON для дедупликации, CSV как журнал и таблица Excel для просмотра."""
from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path

from .config import REPAIR_LABELS
from .scoring import Assessment

CSV_FIELDS = ["date", "id", "item", "title", "price", "market", "below_pct", "sell_price",
              "repair_cost", "margin", "roi_pct", "risk", "risk_reasons", "repairs", "city", "url"]


class SeenStore:
    """id объявления -> цена, при которой мы его уже разбирали."""

    def __init__(self, path: Path | None = None):
        self.path = Path(path) if path else None
        self.data: dict[str, dict] = {}
        if self.path and self.path.exists():
            try:
                self.data = json.loads(self.path.read_text(encoding="utf-8"))
            except (OSError, ValueError):
                self.data = {}

    def needs_review(self, listing_id: str, price: int) -> bool:
        """Новое объявление или цена упала с прошлого раза."""
        prev = self.data.get(listing_id)
        return prev is None or price < prev["price"]

    def mark(self, listing_id: str, price: int, verdict: str):
        self.data[listing_id] = {"price": price, "verdict": verdict,
                                 "t": datetime.now().isoformat(timespec="seconds")}

    def save(self):
        if self.path:
            self.path.parent.mkdir(parents=True, exist_ok=True)
            self.path.write_text(json.dumps(self.data, ensure_ascii=False), encoding="utf-8")


class ResultsLog:
    """Найденные варианты. CSV (utf-8-sig) хранит всё, а data/results.xlsx пересобирается из него
    после каждой находки: лучшие сверху, ссылки кликабельны."""

    XLSX_COLUMNS = [  # (заголовок, ключ в CSV, ширина)
        ("Дата", "date", 17), ("Модель", "item", 26), ("Объявление", "title", 40),
        ("Цена, ₽", "price", 10), ("Рынок (медиана), ₽", "market", 14), ("Ниже рынка, %", "below_pct", 12),
        ("Продажа, ₽", "sell_price", 11), ("Ремонт, ₽", "repair_cost", 10), ("Прибыль, ₽", "margin", 11),
        ("Окупаемость, %", "roi_pct", 13), ("Риск, из 100", "risk", 11), ("Почему такой риск", "risk_reasons", 50),
        ("Что чинить", "repairs", 18), ("Город", "city", 14), ("Ссылка", "url", 12),
    ]
    NUMERIC = {"price", "market", "below_pct", "sell_price", "repair_cost", "margin", "roi_pct", "risk"}

    def __init__(self, path: Path):
        self.path = Path(path)

    @property
    def xlsx_path(self) -> Path:
        return self.path.with_suffix(".xlsx")

    def append(self, a: Assessment):
        self.path.parent.mkdir(parents=True, exist_ok=True)
        new = not self.path.exists()
        with open(self.path, "a", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=CSV_FIELDS)
            if new:
                w.writeheader()
            w.writerow({
                "date": datetime.now().isoformat(timespec="seconds"), "id": a.listing.id,
                "item": a.group, "title": a.listing.title, "price": a.listing.price,
                "market": a.market, "below_pct": round(a.below_pct), "sell_price": a.sell_price,
                "repair_cost": a.repair_cost, "margin": a.margin, "roi_pct": round(a.roi_pct),
                "risk": a.risk, "risk_reasons": "; ".join(a.risk_reasons),
                "repairs": ",".join(a.repairs), "city": a.listing.city, "url": a.listing.url,
            })
        self.export_xlsx()

    def rows(self) -> list[dict]:
        if not self.path.exists():
            return []
        with open(self.path, "r", encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def export_xlsx(self) -> bool:
        """Пересобрать таблицу Excel. False, если файл не удалось записать (обычно он открыт в Excel)."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
        from openpyxl.utils import get_column_letter

        latest: dict[str, dict] = {}
        for r in self.rows():                      # одно объявление: последняя запись (цена могла упасть)
            latest[r["id"]] = r
        rows = sorted(latest.values(), key=lambda r: int(r["margin"] or 0), reverse=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Варианты"
        ws.append([title for title, _, _ in self.XLSX_COLUMNS])
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for r in rows:
            values = []
            for _, key, _ in self.XLSX_COLUMNS:
                v = r.get(key, "")
                if key in self.NUMERIC:
                    try:
                        v = int(float(v))
                    except ValueError:
                        pass
                elif key == "repairs":
                    v = ", ".join(REPAIR_LABELS.get(x, x) for x in v.split(",") if x)
                elif key == "url":
                    v = "Открыть"
                values.append(v)
            ws.append(values)
            link = ws.cell(row=ws.max_row, column=len(self.XLSX_COLUMNS))
            link.hyperlink = r.get("url", "")
            link.font = Font(color="0563C1", underline="single")
        for n, (_, key, width) in enumerate(self.XLSX_COLUMNS, 1):
            ws.column_dimensions[get_column_letter(n)].width = width
            if key in self.NUMERIC and key != "risk":
                for row in ws.iter_rows(min_row=2, min_col=n, max_col=n):
                    row[0].number_format = "# ##0"
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        try:
            self.xlsx_path.parent.mkdir(parents=True, exist_ok=True)
            tmp = self.xlsx_path.with_name(self.xlsx_path.stem + ".tmp.xlsx")
            wb.save(tmp)
            tmp.replace(self.xlsx_path)
            return True
        except OSError as exc:
            import logging
            logging.getLogger("resell_radar.storage").warning(
                "Не удалось записать %s (закройте файл в Excel, он обновится при следующей находке): %s",
                self.xlsx_path, exc)
            return False
